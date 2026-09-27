import os, sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from app import app
from models import Faculty, Subject

def run_checks():
    with app.app_context():
        fac = Faculty.query.filter_by(status='Active').first()
        assert fac is not None, "No active faculty found in database"
        fac_id = fac.id

    with app.test_client() as client:
        with client.session_transaction() as sess:
            sess['faculty_id'] = fac_id
            sess['faculty_name'] = fac.full_name

        resp = client.get('/faculty/marks?academic_year=2026-27&subject_id=10&division=A&semester=1')
        print('Status code:', resp.status_code)
        assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"

        html = resp.get_data(as_text=True)

        assert 'Viva / Interview' in html, "Missing 'Viva / Interview' in page"
        assert 'Coding / Practical' in html, "Missing 'Coding / Practical' in page"
        assert 'Journal' in html, "Missing 'Journal' in page"
        print("SUCCESS: Found 'Viva / Interview', 'Coding / Practical', 'Journal' in /faculty/marks HTML!")

        # Verify Excel and PDF export logic directly with practical subject
        from services.export_marks_service import generate_internal_marks_excel, generate_internal_marks_pdf
        from openpyxl import load_workbook
        import io

        with app.app_context():
            sub = Subject.query.filter_by(id=10).first()
            data_dict = {
                "subject": {
                    "id": sub.id,
                    "subject_code": sub.subject_code,
                    "subject_name": sub.subject_name,
                    "subject_type": sub.subject_type,
                    "internal_marks": sub.internal_marks or 25,
                    "maxes": {"internal_exam": 10.0, "practical_eval": 10.0, "journal": 5.0}
                },
                "semester": 1,
                "division": "A",
                "academic_year": "2026-27",
                "students": [
                    {
                        "roll_number": 16,
                        "enrollment_no": "BCA15226000021",
                        "full_name": "akash raval",
                        "breakdown": {"internal_exam": 8.5, "practical_eval": 9.0, "journal": 4.5},
                        "marks_obtained": 22.0
                    }
                ]
            }

            excel_buf = generate_internal_marks_excel(data_dict, "Test Faculty")
            wb = load_workbook(excel_buf)
            ws = wb.active
            headers = [cell.value for cell in ws[7] if cell.value is not None]
            print("Exported Excel Headers:", headers)
            assert any("Viva / Interview" in str(h) for h in headers), "Viva / Interview not found in Excel headers"
            assert any("Coding / Practical" in str(h) for h in headers), "Coding / Practical not found in Excel headers"
            assert any("Journal" in str(h) for h in headers), "Journal not found in Excel headers"
            print("SUCCESS: Excel export contains 'Viva / Interview' and 'Coding / Practical'!")

            pdf_buf = generate_internal_marks_pdf(data_dict, "Test Faculty")
            pdf_bytes = pdf_buf.getvalue() if hasattr(pdf_buf, 'getvalue') else pdf_buf
            assert len(pdf_bytes) > 1000, "PDF export produced empty or invalid output"
            print("SUCCESS: PDF export generated valid PDF binary!")

        print("ALL PRACTICAL COMPONENT CHECKS PASSED SUCCESSFULLY!")

if __name__ == '__main__':
    run_checks()
